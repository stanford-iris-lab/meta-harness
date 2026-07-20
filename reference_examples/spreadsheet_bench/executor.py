"""Subprocess code executor for the spreadsheet agent.

The persistent task state is the `.xlsx` file on disk. Each execution runs a block of
model-generated Python in a fresh working directory that contains `input.xlsx` (a copy
of the given input). By convention the code reads `input.xlsx` from the current
directory and, for a *solution*, writes `output.xlsx` to the current directory.

This is deliberately simple (no Docker / no persistent kernel): re-running the final
solution on each OJ test-case input reproduces the program's effect. Must be run inside
a Slurm job, never the login node.
"""

from __future__ import annotations

import os
import shutil
import subprocess
import sys
from dataclasses import dataclass
from pathlib import Path

# numpy/OpenBLAS default to one thread per core, which on a busy shared node both wastes
# memory and emits "blas_thread_init: pthread_create failed" on stderr — noise that would
# otherwise be fed back to the model as part of the execution observation. Spreadsheet
# work is not BLAS-bound, so pin every math backend to a single thread.
_THREAD_ENV = {
    "OPENBLAS_NUM_THREADS": "1",
    "OMP_NUM_THREADS": "1",
    "MKL_NUM_THREADS": "1",
    "NUMEXPR_NUM_THREADS": "1",
    "VECLIB_MAXIMUM_THREADS": "1",
}

INPUT_NAME = "input.xlsx"
OUTPUT_NAME = "output.xlsx"
SNIPPET_NAME = "_snippet.py"
MAX_STREAM_CHARS = 8000


@dataclass
class ExecResult:
    ok: bool  # process exited 0 AND (output produced when expected)
    stdout: str
    stderr: str
    output_path: Path | None
    timed_out: bool = False
    returncode: int | None = None


def _truncate(text: str, limit: int = MAX_STREAM_CHARS) -> str:
    if text is None:
        return ""
    if len(text) <= limit:
        return text
    head = text[: limit // 2]
    tail = text[-limit // 2 :]
    return f"{head}\n... [truncated {len(text) - limit} chars] ...\n{tail}"


def run_code(
    code: str,
    input_path: str | Path,
    workdir: str | Path,
    timeout: int = 60,
    expect_output: bool = True,
    python_exe: str | None = None,
) -> ExecResult:
    """Run `code` in `workdir` with a fresh `input.xlsx`; capture stdout/stderr.

    If `expect_output`, `ok` also requires `output.xlsx` to have been written.
    """
    workdir = Path(workdir)
    workdir.mkdir(parents=True, exist_ok=True)

    shutil.copyfile(input_path, workdir / INPUT_NAME)
    out_path = workdir / OUTPUT_NAME
    if out_path.exists():
        out_path.unlink()

    (workdir / SNIPPET_NAME).write_text(code)
    py = python_exe or sys.executable

    timed_out = False
    returncode: int | None = None
    try:
        proc = subprocess.run(
            [py, SNIPPET_NAME],
            cwd=str(workdir),
            timeout=timeout,
            capture_output=True,
            text=True,
            env={**os.environ, **_THREAD_ENV},
        )
        stdout, stderr, returncode = proc.stdout, proc.stderr, proc.returncode
    except subprocess.TimeoutExpired as e:
        stdout = e.stdout or ""
        if isinstance(stdout, bytes):
            stdout = stdout.decode("utf-8", "replace")
        stderr = (e.stderr or "")
        if isinstance(stderr, bytes):
            stderr = stderr.decode("utf-8", "replace")
        stderr += f"\n[TIMEOUT after {timeout}s]"
        timed_out = True

    output_path = out_path if out_path.exists() else None
    proc_ok = (returncode == 0) and not timed_out
    ok = proc_ok and (output_path is not None if expect_output else True)
    return ExecResult(
        ok=ok,
        stdout=_truncate(stdout),
        stderr=_truncate(stderr),
        output_path=output_path,
        timed_out=timed_out,
        returncode=returncode,
    )


def _self_test() -> None:
    """Smoke check: write a tiny xlsx, run a snippet that copies a value across."""
    import tempfile

    import openpyxl

    tmp = Path(tempfile.mkdtemp(prefix="sb_exec_"))
    src = tmp / "src.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = 21
    wb.save(src)

    code = (
        "import openpyxl\n"
        "wb = openpyxl.load_workbook('input.xlsx')\n"
        "ws = wb.active\n"
        "ws['B1'] = ws['A1'].value * 2\n"
        "print('doubled', ws['B1'].value)\n"
        "wb.save('output.xlsx')\n"
    )
    res = run_code(code, src, tmp / "work", timeout=30)
    assert res.ok, f"exec failed: {res.stderr}"
    got = openpyxl.load_workbook(res.output_path, data_only=True).active["B1"].value
    assert got == 42, f"expected 42, got {got}"
    print("executor self-test OK:", res.stdout.strip())
    shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    _self_test()
