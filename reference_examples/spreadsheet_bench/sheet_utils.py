"""openpyxl helpers: sheet preview, answer-position parsing, and cell comparison.

The comparison logic mirrors SpreadsheetBench's `compare_cell_value`:
  - numeric values rounded to 2 decimals (within a small tolerance)
  - datetime -> Excel serial float; time -> string
  - numeric strings compared as numbers
  - empty string and None are equivalent
  - type mismatch -> not equal
Workbooks are always loaded with data_only=True so we read cached/literal values.
"""

from __future__ import annotations

from datetime import date, datetime, time

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

_EXCEL_EPOCH = datetime(1899, 12, 30)


def _excel_serial(dt: datetime) -> float:
    delta = dt - _EXCEL_EPOCH
    return delta.days + delta.seconds / 86400.0


def _to_float(v):
    """Return (float_value, is_numeric)."""
    if isinstance(v, bool):
        return float(v), True
    if isinstance(v, (int, float)):
        return float(v), True
    if isinstance(v, str):
        s = v.strip().replace(",", "")
        if s.endswith("%"):
            try:
                return float(s[:-1]) / 100.0, True
            except ValueError:
                return 0.0, False
        try:
            return float(s), True
        except ValueError:
            return 0.0, False
    return 0.0, False


def compare_cell_value(v1, v2, tol: float = 1e-6) -> bool:
    """True if two cell values are equivalent under SpreadsheetBench semantics."""
    empty1 = v1 is None or v1 == ""
    empty2 = v2 is None or v2 == ""
    if empty1 or empty2:
        return empty1 and empty2

    if isinstance(v1, datetime):
        v1 = _excel_serial(v1)
    elif isinstance(v1, date):
        v1 = _excel_serial(datetime(v1.year, v1.month, v1.day))
    if isinstance(v2, datetime):
        v2 = _excel_serial(v2)
    elif isinstance(v2, date):
        v2 = _excel_serial(datetime(v2.year, v2.month, v2.day))

    if isinstance(v1, time):
        v1 = str(v1)
    if isinstance(v2, time):
        v2 = str(v2)

    n1, ok1 = _to_float(v1)
    n2, ok2 = _to_float(v2)
    if ok1 and ok2:
        return abs(round(n1, 2) - round(n2, 2)) <= tol
    if ok1 != ok2:
        return False
    return str(v1).strip() == str(v2).strip()


def _split_positions(answer_position: str) -> list[str]:
    return [p for p in (answer_position or "").split(",") if p.strip()]


def parse_answer_position(answer_position: str) -> list[tuple[str | None, str]]:
    """Parse into a list of (sheet_name_or_None, range_str)."""
    out: list[tuple[str | None, str]] = []
    for part in _split_positions(answer_position):
        part = part.strip()
        if "!" in part:
            sheet, rng = part.rsplit("!", 1)
            sheet = sheet.strip().strip("'").strip('"')
        else:
            sheet, rng = None, part
        out.append((sheet or None, rng.strip()))
    return out


def cells_in_range(range_str: str) -> list[str]:
    """All cell addresses in an A1-style range (single cell allowed)."""
    min_col, min_row, max_col, max_row = range_boundaries(range_str)
    coords = []
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            coords.append(f"{get_column_letter(c)}{r}")
    return coords


def _get_ws(wb, sheet_name):
    if sheet_name is None:
        return wb.active
    if sheet_name in wb.sheetnames:
        return wb[sheet_name]
    return None


def compare_workbooks(
    gt_path, out_path, answer_position: str, tol: float = 1e-6,
    default_sheet: str | None = None,
) -> tuple[bool, str]:
    """Compare produced vs ground-truth at answer positions. Returns (ok, detail).

    `default_sheet` is used for positions with no "Sheet!" qualifier (some releases
    carry the sheet name in a separate `answer_sheet` field). A malformed range never
    raises — it is reported as a mismatch so one bad record can't crash a run.
    """
    gt = load_workbook(gt_path, data_only=True)
    out = load_workbook(out_path, data_only=True)
    positions = parse_answer_position(answer_position)
    if not positions:
        return False, "empty answer_position"

    for sheet, rng in positions:
        sheet = sheet or (default_sheet or None)
        gt_ws = _get_ws(gt, sheet)
        out_ws = _get_ws(out, sheet)
        if gt_ws is None:
            return False, f"gt sheet missing: {sheet}"
        if out_ws is None:
            return False, f"output sheet missing: {sheet}"
        try:
            coords = cells_in_range(rng)
        except ValueError:
            return False, f"unparseable range: {rng!r}"
        for coord in coords:
            v1 = gt_ws[coord].value
            v2 = out_ws[coord].value
            if not compare_cell_value(v1, v2, tol):
                loc = f"{sheet + '!' if sheet else ''}{coord}"
                return False, f"mismatch at {loc}: gt={v1!r} out={v2!r}"
    return True, "ok"


def preview_spreadsheet(path, max_rows: int = 20, max_cols: int = 20) -> str:
    """Human-readable preview of the first rows of each sheet (values, data_only)."""
    wb = load_workbook(path, data_only=True)
    blocks = []
    for ws in wb.worksheets:
        dims = f"{ws.max_row} rows x {ws.max_column} cols"
        lines = [f"### Sheet: {ws.title} ({dims})"]
        n_cols = min(ws.max_column, max_cols)
        header = ["    "] + [get_column_letter(c) for c in range(1, n_cols + 1)]
        lines.append(" | ".join(header))
        for r in range(1, min(ws.max_row, max_rows) + 1):
            row_cells = []
            for c in range(1, n_cols + 1):
                val = ws.cell(row=r, column=c).value
                row_cells.append("" if val is None else str(val))
            lines.append(" | ".join([f"r{r}"] + row_cells))
        if ws.max_row > max_rows or ws.max_column > max_cols:
            lines.append(
                f"... (truncated to {max_rows} rows x {max_cols} cols) ..."
            )
        blocks.append("\n".join(lines))
    return "\n\n".join(blocks)
