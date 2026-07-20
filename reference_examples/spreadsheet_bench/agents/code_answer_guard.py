"""Deterministic "value-not-code" answer guard: reject source/macro/formula text written
into answer cells and force ONE targeted re-solve.

Parent: ``baseline_react`` (a bare ``SpreadsheetAgent``). The baseline accepts the FIRST
program that merely runs (``if res.ok: break``) and never inspects what was written. On the dev
set the single most legible failure is a *how-to misread*: when an instruction is phrased as a
question — "how can I write / optimize the VBA / a macro / a formula to do X" — the small model
answers the QUESTION. It writes macro source and a step-by-step explanation into the answer
cells instead of PERFORMING X on the data and writing the resulting literal values. The program
runs, so the baseline keeps it and it scores 0 on every case. The base prompt already says
"write literal computed values, NOT formulas" and the model still does this: under the how-to
framing a passive upfront directive is exactly the lever that fails, which is why appending yet
another instruction does not fix it — the correction has to confront the concrete emitted code.

ONE new mechanism: a **deterministic** post-solve *content* guard. After a solution runs, the
scaffold reads the produced output (``data_only=False`` so formula strings are visible) and
classifies each written answer cell *by its form*. If any answer cell holds an artifact that can
never be a computed data value — an Excel formula (``=...``), recognisable source code / a macro
(BASIC/VBA or a couple of cross-language structural tokens), or a multi-word phrase that *names*
code ("...VBA...", "...subroutine...") — the loop withholds acceptance ONCE, feeds back the
specific offending cell(s) plus the general principle (write the VALUES produced by performing
the operation, never code/prose about it), and requests a corrected solve. A value-clean output
fires nothing: no extra LLM call, byte-for-byte the baseline. That property is deliberate — it is
why the guard cannot regress a currently-passing task (both passers emit ordinary filtered/
aggregated data values) and why it adds no per-call infra-flake exposure on tasks that pass.

Distinct from the two tested post-solve candidates:

  * ``answer_type_gate`` decided a violation from the *column's* numeric dominance in the
    non-answer rows, so it went structurally **blind** exactly here — when the answer range spans
    the whole data column (a text header + blank tail leave <4 typeable cells) or the intended
    value is itself text, it flags nothing and macro prose sails through. This guard recognises
    the *written artifact directly* (column-type independent), so it fires precisely on the
    misread the type gate misses.
  * ``readback_verify`` showed the values back and let the *model* decide pass/revise (it
    rubber-stamped its own output and tied). Here the reject/accept decision is a rule, not the
    model's opinion, so it cannot collude.

Recognising "an answer cell holds a program/macro/formula rather than a data value" is a general
spreadsheet-answer validity property (every such task wants values, not code, in answer cells);
the tokens matched are general programming/office markers, never a dataset's content — no
hardcoded column, row, value, or task-specific rule. The first runnable program is always kept
as the fallback, and the guard only fires on an output that is *already* certain to be wrong, so
it can only replace it with a later runnable solve — never do worse than the baseline would.
"""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

from agent import PASSTHROUGH_SOLUTION, SpreadsheetAgent
from executor import run_code
from sheet_utils import parse_answer_position

# Structural signatures of source code / a spreadsheet macro. These are general
# programming/office tokens (BASIC/VBA + a couple of cross-language markers), each carrying
# code punctuation or a reserved keyword that essentially never occurs inside a literal data
# value or a short plain-language answer. Matching one means the cell holds *code about the
# task*, not a task value.
_CODE_PATTERNS = [
    re.compile(r"^\s*(private\s+|public\s+)?(sub|function)\s+\w+\s*\(", re.I),
    re.compile(r"\bdim\s+\w+\s+as\b", re.I),
    re.compile(r"\bcreateobject\s*\(", re.I),
    re.compile(
        r"\bapplication\.(screenupdating|calculation|enableevents|worksheetfunction|displayalerts)\b",
        re.I,
    ),
    re.compile(r"\bfor\s+\w+\s*=.*\bto\b", re.I),
    re.compile(r"\bfor\s+each\b", re.I),
    re.compile(r"\b(range|cells|worksheets|activesheet|activeworkbook)\s*\(", re.I),
    re.compile(r"\b(msgbox|xlup|xldown)\b", re.I),
    re.compile(r"\bscripting\.dictionary\b", re.I),
    re.compile(r"\bxlcalculation\w*\b", re.I),
    re.compile(r"\bend\s+(sub|function|if|with)\b", re.I),
    re.compile(r"\bdef\s+\w+\s*\(", re.I),
]
# A prose sentence that *names* code (e.g. "Optimized VBA to remove duplicate IDs efficiently"):
# the words 'vba'/'subroutine' inside a multi-word string. Both are almost never a legitimate
# multi-word data value, and the length gate keeps a bare one-word cell from ever matching.
_DESCRIBES_CODE = re.compile(r"\b(vba|subroutine)\b", re.I)


def _resolve_ws(wb, sheet_name):
    if sheet_name is None:
        return wb.active
    if sheet_name in wb.sheetnames:
        return wb[sheet_name]
    return None


def _looks_like_code(value) -> str | None:
    """Return a short reason string if the cell value is code/macro/formula, else None."""
    if not isinstance(value, str):
        return None
    s = value.strip()
    if not s:
        return None
    if s.startswith("="):
        return "an Excel formula (stored without a cached value, so the scorer reads it as empty)"
    for pat in _CODE_PATTERNS:
        if pat.search(s):
            return "source code / a macro"
    if len(s.split()) >= 3 and _DESCRIBES_CODE.search(s):
        return "a description of code rather than a data value"
    return None


class AgentHarness(SpreadsheetAgent):
    # Lifted only to fund the single correction round the guard may request; the mechanism is
    # the deterministic content guard, not the larger budget.
    MAX_TURNS = 7
    MAX_FIX_ROUNDS = 1     # at most one guard-triggered correction per instruction
    MAX_VIOLATIONS = 8     # cap the cells reported back

    # ── deterministic content check ────────────────────────────────────────
    def _code_violations(self, task, output_path) -> list[tuple[str, object, str]]:
        """List (loc, written_value, reason) answer cells that hold code/macro/formula text.

        Reads the produced output with ``data_only=False`` so formula strings are visible. Any
        failure degrades to an empty list → the baseline accept-first behaviour.
        """
        try:
            out_wb = load_workbook(output_path, data_only=False)
        except Exception:
            return []
        default_sheet = getattr(task, "answer_sheet", "") or None
        violations: list[tuple[str, object, str]] = []
        try:
            for sheet, rng in parse_answer_position(task.answer_position):
                out_ws = _resolve_ws(out_wb, sheet or default_sheet)
                if out_ws is None:
                    continue
                try:
                    min_c, min_r, max_c, max_r = range_boundaries(rng)
                except Exception:
                    continue
                if None in (min_c, min_r, max_c, max_r):
                    continue
                for c in range(min_c, max_c + 1):
                    for r in range(min_r, max_r + 1):
                        reason = _looks_like_code(out_ws.cell(row=r, column=c).value)
                        if reason:
                            violations.append(
                                (f"{get_column_letter(c)}{r}", out_ws.cell(row=r, column=c).value, reason)
                            )
                            if len(violations) >= self.MAX_VIOLATIONS:
                                return violations
        except Exception:
            return violations
        return violations

    def _fix_prompt(self, violations) -> str:
        lines = []
        for loc, val, reason in violations:
            shown = repr(val)
            if len(shown) > 70:
                shown = shown[:67] + "...'"
            lines.append(f"  - {loc}: you wrote {shown} — that is {reason}.")
        body = "\n".join(lines)
        return (
            "Automated content check (deterministic — it inspects only the FORM of what your "
            "program wrote into the answer cells, never an answer key). Your program ran, but "
            "these answer cells hold code / a macro / a formula / a description instead of a "
            "computed data value, so they will be scored wrong:\n"
            f"{body}\n\n"
            "This is the signature of a misread: when an instruction is phrased as a question "
            '("how can I write / optimize a macro / a formula / code to do X"), the deliverable '
            "is NOT that code or an explanation of it — it is the literal VALUES that result "
            "from actually PERFORMING X on the data. Do the operation yourself in Python and "
            "write the concrete resulting number/string/date into each answer cell (never code, "
            "a macro, a formula, or prose). Resubmit the complete program:\n"
            "ACTION: solve\n"
            "```python\n"
            "# read input.xlsx, perform the operation, write literal values into the answer "
            "cells, save output.xlsx\n"
            "```"
        )

    # ── main loop (baseline structure + the deterministic content guard) ────
    def run_loop(self, task, primary_input: Path, workdir: Path) -> dict:
        messages = [
            {"role": "system", "content": self.build_system_prompt()},
            {"role": "user", "content": self.build_user_prompt(task, primary_input)},
        ]
        solution_code = None   # first runnable solution kept as the fallback
        produced = False
        turns = 0
        code = None
        fix_used = 0

        for turn in range(self.MAX_TURNS):
            turns = turn + 1
            messages = self.summarize_context(messages)
            response = self.call_llm(messages)
            messages.append({"role": "assistant", "content": response})

            action, code = self.parse_action(response)
            if code is None:
                messages.append({
                    "role": "user",
                    "content": (
                        "No code block found. Respond with an ACTION line and one fenced "
                        "```python``` block."
                    ),
                })
                continue

            run_dir = workdir / f"turn_{turns}"
            if action == "solve":
                res = run_code(
                    code, primary_input, run_dir,
                    timeout=self.CODE_TIMEOUT, expect_output=True,
                    python_exe=self.python_exe,
                )
                if res.ok:
                    solution_code = code
                    produced = True
                    # Deterministic guard: withhold acceptance only for an objective code/
                    # macro/formula artifact in an answer cell, and only once. A value-clean
                    # output is accepted immediately with no extra LLM call (identical to the
                    # baseline).
                    if fix_used < self.MAX_FIX_ROUNDS:
                        violations = self._code_violations(task, res.output_path)
                        if violations:
                            messages.append(
                                {"role": "user", "content": self._fix_prompt(violations)}
                            )
                            fix_used += 1
                            continue
                    break  # accept: value-clean, or the one correction round is spent
                messages.append({"role": "user", "content": self.observe(res, "solve")})
            else:  # explore
                res = run_code(
                    code, primary_input, run_dir,
                    timeout=self.CODE_TIMEOUT, expect_output=False,
                    python_exe=self.python_exe,
                )
                messages.append({"role": "user", "content": self.observe(res, "explore")})

        if solution_code is None:
            solution_code = code if code else PASSTHROUGH_SOLUTION
        return {
            "solution_code": solution_code,
            "n_turns": turns,
            "produced": produced,
            "trajectory": messages,
        }
