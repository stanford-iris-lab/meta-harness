"""Read-back self-verification gate before acceptance.

The baseline ReAct loop accepts the FIRST solution that merely *runs*
(`if res.ok: break`) — it never looks at what it actually wrote, so a program that
executes cleanly but computes the wrong values (or writes prose/formulas into
numeric answer cells) is accepted as-is. The dominant failure mode on the dev set
is exactly this "accepted-but-wrong".

ONE mechanism is added here: after a solution runs successfully, do NOT accept it
immediately. Instead re-read the values the program actually wrote into the
`answer_position` cells (shown next to the original input values) and make the model
reconcile them against the instruction before accepting. The model must reply
`VERIFY: pass` (accept) or `VERIFY: revise` + a corrected program. The revise
criterion is deliberately conservative (name a concrete offending cell) so a correct
solution is not second-guessed into a worse one. No ground truth is ever used — the
gate reads only the produced `output.xlsx`, the input, and the instruction, which is
the self-check the scaffold otherwise lacks.

Everything else (exploration, parsing, observations) is identical to the baseline.
"""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from agent import PASSTHROUGH_SOLUTION, SpreadsheetAgent
from executor import run_code
from sheet_utils import cells_in_range, parse_answer_position

_VERDICT_RE = re.compile(r"VERIFY:\s*(pass|revise)", re.IGNORECASE)


class AgentHarness(SpreadsheetAgent):
    # The verification gate needs a couple of turns beyond the usual explore+solve to
    # actually run; MAX_TURNS is lifted only to fund those rounds (the mechanism is the
    # read-back gate, not the larger budget).
    MAX_TURNS = 8
    VERIFY_ROUNDS = 2      # max read-back / revise cycles per runnable solution
    MAX_VERIFY_CELLS = 60  # cap answer cells shown back (keeps context bounded)

    # ── read-back helpers ──────────────────────────────────────────────────
    def _read_answer_cells(self, path, answer_position, default_sheet):
        """List (loc, value) for the answer cells in `path`. None if unreadable."""
        try:
            wb = load_workbook(path, data_only=True)
        except Exception:
            return None
        rows: list[tuple[str, object]] = []
        for sheet, rng in parse_answer_position(answer_position):
            sheet = sheet or default_sheet
            if sheet is None:
                ws = wb.active
            elif sheet in wb.sheetnames:
                ws = wb[sheet]
            else:
                continue
            try:
                coords = cells_in_range(rng)
            except Exception:
                continue
            prefix = f"{sheet}!" if sheet else ""
            for coord in coords:
                rows.append((f"{prefix}{coord}", ws[coord].value))
                if len(rows) >= self.MAX_VERIFY_CELLS:
                    return rows
        return rows

    def _build_verify_prompt(self, task, input_path, output_path):
        """Verify prompt showing (original -> written) answer cells. None to skip."""
        default_sheet = getattr(task, "answer_sheet", "") or None
        written = self._read_answer_cells(output_path, task.answer_position, default_sheet)
        if not written:
            return None  # nothing to read back → fall back to base (accept)
        original = dict(self._read_answer_cells(input_path, task.answer_position,
                                                default_sheet) or [])
        lines = [f"{loc}: original={original.get(loc, '')!r} -> written={val!r}"
                 for loc, val in written]
        note = (" (first cells only; the rest follow the same pattern)"
                if len(written) >= self.MAX_VERIFY_CELLS else "")
        readback = "\n".join(lines)
        return (
            "Your program ran and wrote output.xlsx. Do NOT assume it is correct — "
            "verify it before it is accepted.\n\n"
            f"These are the values your program actually wrote into the answer cells{note}:\n"
            f"{readback}\n\n"
            "First, using ONLY the instruction, restate in one or two sentences what these "
            "answer cells should contain: the expected data TYPE (number / text / date), how "
            "many cells should be populated and in what shape, the transformation rule, and any "
            "convention the instruction implies (for example first-vs-last match, case "
            "sensitivity, rounding, how blank or boundary rows are handled, and that values "
            "must be literals not formulas — a formula reads back as None here).\n"
            "Then compare each written value above to that expectation.\n\n"
            "Reply with exactly one verdict on the first line:\n"
            "  VERIFY: pass    — the written values satisfy the instruction.\n"
            "  VERIFY: revise  — you can name a SPECIFIC cell whose value violates it.\n"
            "Only choose revise when you can point to a concrete cell and the concrete reason "
            "it is wrong; if the output plausibly satisfies the instruction, choose pass. If you "
            "revise, follow the verdict line with a corrected, self-contained program:\n"
            "ACTION: solve\n"
            "```python\n"
            "# full program: read input.xlsx, compute, write literal values, save output.xlsx\n"
            "```"
        )

    def _parse_verdict(self, response):
        m = _VERDICT_RE.search(response or "")
        return m.group(1).lower() if m else None

    def _maybe_request_verify(self, task, input_path, output_path, messages, verify_used):
        """Append a verify prompt; return True if a verification round should run."""
        if verify_used >= self.VERIFY_ROUNDS:
            return False
        prompt = self._build_verify_prompt(task, input_path, output_path)
        if prompt is None:
            return False
        messages.append({"role": "user", "content": prompt})
        return True

    # ── main loop (baseline structure + the verification gate) ─────────────
    def run_loop(self, task, primary_input: Path, workdir: Path) -> dict:
        messages = [
            {"role": "system", "content": self.build_system_prompt()},
            {"role": "user", "content": self.build_user_prompt(task, primary_input)},
        ]
        solution_code = None  # best runnable solution so far (fallback if verify stalls)
        produced = False
        turns = 0
        code = None
        verify_used = 0
        awaiting_verify = False

        for turn in range(self.MAX_TURNS):
            turns = turn + 1
            messages = self.summarize_context(messages)
            response = self.call_llm(messages)
            messages.append({"role": "assistant", "content": response})

            # ── handle a response to a verification request ────────────────
            if awaiting_verify:
                awaiting_verify = False
                verdict = self._parse_verdict(response)
                _, new_code = self.parse_action(response)
                if verdict == "revise" and new_code is not None:
                    res = run_code(
                        new_code, primary_input, workdir / f"turn_{turns}",
                        timeout=self.CODE_TIMEOUT, expect_output=True,
                        python_exe=self.python_exe,
                    )
                    if res.ok:
                        solution_code = new_code  # adopt the corrected program
                        produced = True
                        if self._maybe_request_verify(
                            task, primary_input, res.output_path, messages, verify_used
                        ):
                            verify_used += 1
                            awaiting_verify = True
                            continue
                        break  # verify budget spent → accept the corrected program
                    # corrected program broke: keep prior solution, feed back the error
                    messages.append({"role": "user", "content": self.observe(res, "solve")})
                    continue
                break  # VERIFY: pass (or no actionable revision) → accept best-so-far

            # ── normal explore / solve turn ────────────────────────────────
            action, code = self.parse_action(response)
            if code is None:
                messages.append({
                    "role": "user",
                    "content": (
                        "No code block found. Respond with an ACTION line and one fenced "
                        "```python``` block."
                    ),
                })
                continue

            run_dir = workdir / f"turn_{turns}"
            if action == "solve":
                res = run_code(
                    code, primary_input, run_dir,
                    timeout=self.CODE_TIMEOUT, expect_output=True,
                    python_exe=self.python_exe,
                )
                if res.ok:
                    solution_code = code
                    produced = True
                    if self._maybe_request_verify(
                        task, primary_input, res.output_path, messages, verify_used
                    ):
                        verify_used += 1
                        awaiting_verify = True
                        continue  # do not accept yet — verify what was written
                    break  # no verify budget/read-back → base behavior (accept)
                messages.append({"role": "user", "content": self.observe(res, "solve")})
            else:  # explore
                res = run_code(
                    code, primary_input, run_dir,
                    timeout=self.CODE_TIMEOUT, expect_output=False,
                    python_exe=self.python_exe,
                )
                messages.append({"role": "user", "content": self.observe(res, "explore")})

        if solution_code is None:
            solution_code = code if code else PASSTHROUGH_SOLUTION
        return {
            "solution_code": solution_code,
            "n_turns": turns,
            "produced": produced,
            "trajectory": messages,
        }
